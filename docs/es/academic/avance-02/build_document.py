"""Build the academic DOCX from the reviewed Markdown source."""
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[4]
SOURCE = Path(__file__).parent
DEST = ROOT / "entregables/Avance2_ORVIX_2026-10-11/Documentacion_ORVIX_Avance2.docx"


def element(name, **attrs):
    result = OxmlElement(name)
    for key, value in attrs.items():
        result.set(qn(key), str(value))
    return result


def hyperlink(p, url, label="Consultar fuente"):
    link = element("w:hyperlink", **{"r:id": p.part.relate_to(url, RT.HYPERLINK, is_external=True)})
    run = element("w:r")
    props = element("w:rPr")
    props.append(element("w:color", **{"w:val": "23577A"}))
    props.append(element("w:u", **{"w:val": "single"}))
    run.append(props)
    text = element("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    p._p.append(link)


def inline(p, text):
    for part in re.split(r"(\*\*.*?\*\*|https?://\S+)", text):
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith(("https://", "http://")):
            hyperlink(p, part)
        elif part:
            p.add_run(part)


def set_cell_padding(cell):
    props = cell._tc.get_or_add_tcPr()
    margins = element("w:tcMar")
    for name in ("top", "left", "bottom", "right"):
        margins.append(element("w:" + name, **{"w:w": 95, "w:type": "dxa"}))
    props.append(margins)


def add_table(doc, data):
    ncols = len(data[0])
    if ncols == 4:
        widths = [0.72, 2.02, 0.60, 3.56]
    elif ncols == 2:
        widths = [1.65, 5.25]
    elif data[0][0].startswith("Tabla"):
        widths = [2.52, 0.72, 3.66]
    else:
        widths = [1.60, 2.45, 2.85]
    table = doc.add_table(rows=1, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    borders = element("w:tblBorders")
    for border in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(element("w:" + border, **{"w:val": "single", "w:sz": 4, "w:color": "D9D9D9"}))
    table._tbl.tblPr.append(borders)
    for row_index, values in enumerate(data):
        row = table.rows[0] if row_index == 0 else table.add_row()
        row._tr.get_or_add_trPr().append(element("w:cantSplit"))
        if row_index == 0:
            row._tr.get_or_add_trPr().append(element("w:tblHeader"))
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_padding(cell)
            cell._tc.get_or_add_tcPr().append(element("w:shd", **{"w:fill": "273B4D" if row_index == 0 else ("F2F5F7" if row_index % 2 == 0 else "FFFFFF")}))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.04
            if (ncols == 4 and col_index in (0, 2)) or (data[0][1] == "Tipo" and col_index == 1):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(p, value)
            for run in p.runs:
                run.font.size = Pt(10.0)
                run.font.color.rgb = RGBColor.from_string("FFFFFF" if row_index == 0 else "000000")
                run.bold = row_index == 0 or run.bold
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 1
    spacer.add_run().font.size = Pt(3)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.70)
    section.bottom_margin = Inches(0.70)
    section.left_margin = Inches(0.80)
    section.right_margin = Inches(0.80)
    section.footer_distance = Inches(0.32)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11.2)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for style_name, size in (("Title", 28), ("Subtitle", 16), ("Heading 1", 18), ("Heading 2", 13.2), ("Heading 3", 11.5)):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_after = Pt(9)
        style.paragraph_format.space_before = Pt(8 if style_name.startswith("Heading") else 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.08
        props = style.element.get_or_add_rPr()
        color = props.find(qn("w:color"))
        if color is not None:
            for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
                color.attrib.pop(qn(attr), None)
        pprops = style.element.find(qn("w:pPr"))
        if pprops is not None:
            for child in list(pprops):
                if child.tag == qn("w:pBdr"):
                    pprops.remove(child)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("ORVIX   ")
    run.font.size = Pt(8.5)
    field = element("w:fldSimple", **{"w:instr": "PAGE"})
    footer._p.append(field)
    doc.core_properties.title = "ORVIX Avance 2 Semana 10"
    doc.core_properties.subject = "Diseño académico y estado técnico de ORVIX"
    doc.core_properties.author = "Andre Alexander Alvarado Lira; Jose Eduardo Alvarado Mejia; Javier Gomez Correa"
    doc.core_properties.keywords = "ORVIX, Avance 2, captura nativa, memoria compartida, requisitos"
    chunks = (SOURCE / "avance-02.md").read_text(encoding="utf-8").split("<!-- page -->")
    for page_index, chunk in enumerate(chunks):
        first_paragraph = True
        lines = chunk.strip().splitlines()
        i = 0
        while i < len(lines):
            text = lines[i].strip()
            if not text:
                i += 1
                continue
            if text.startswith("|"):
                data = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-+:?", c) for c in row):
                        data.append(row)
                    i += 1
                add_table(doc, data)
                continue
            image = re.fullmatch(r"!\[(.+)\]\((.+)\)", text)
            if image:
                caption, relative = image.groups()
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.keep_with_next = True
                picture = p.add_run().add_picture(str(SOURCE / relative), width=Inches(6.9))
                picture._inline.docPr.set("descr", caption)
                caption_p = doc.add_paragraph(caption)
                caption_p.paragraph_format.space_after = Pt(10)
                for r in caption_p.runs:
                    r.font.size = Pt(9.5)
                    r.italic = True
                i += 1
                continue
            if text.startswith("# "):
                p = doc.add_paragraph(text[2:], style="Title" if page_index == 0 else "Heading 1")
            elif text.startswith("## "):
                p = doc.add_paragraph(text[3:], style="Subtitle" if page_index == 0 else "Heading 2")
            else:
                p = doc.add_paragraph()
                inline(p, text)
            if first_paragraph and page_index:
                p.paragraph_format.page_break_before = True
            first_paragraph = False
            if page_index == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(14)
            i += 1
        if page_index == 0:
            doc.paragraphs[0].paragraph_format.space_before = Pt(78)
    DEST.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DEST)
    print(str(DEST))
    print("Planned source pages:", len(chunks))


if __name__ == "__main__":
    main()
