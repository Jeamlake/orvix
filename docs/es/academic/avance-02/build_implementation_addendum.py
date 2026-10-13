"""Author the implementation addendum without overwriting prior deliverables."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[4]
target = root / "entregables/Avance2_ORVIX_2026-10-11/Anexo_Implementacion_ORVIX_Avance2.docx"
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.75)
section.left_margin = section.right_margin = Inches(0.8)
for name, size in (("Normal", 11), ("Title", 23), ("Subtitle", 12), ("Heading 1", 15), ("Heading 2", 12)):
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.08
    # Override inherited title borders as well as direct borders. Word can
    # otherwise add a blue title rule from a retained style/template.
    pr = style.element.get_or_add_pPr()
    for old in list(pr.findall(qn("w:pBdr"))):
        pr.remove(old)
    borders = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right", "between", "bar"):
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:val"), "nil")
        borders.append(edge)
    pr.append(borders)
doc.core_properties.title = "Implementación del avance dos de ORVIX"
doc.core_properties.author = "Equipo ORVIX"
doc.core_properties.subject = "Interfaz RAW, instalador y validación de ORVIX 0.2.1"


def para(text, style=None):
    return doc.add_paragraph(text, style)


def heading(text, new_page=False):
    p = doc.add_heading(text, 1)
    p.paragraph_format.page_break_before = new_page
    return p


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [3.2, 1.2, 2.5]
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:" + side)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            node.set(qn("w:" + key), value)
        borders.append(node)
    t._tbl.tblPr.append(borders)
    for i, cell in enumerate(t.rows[0].cells):
        cell.text = headers[i]
    for values in rows:
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "253B53" if ri == 0 else ("F3F5F7" if ri % 2 == 0 else "FFFFFF"))
            pr.append(shade)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                node = OxmlElement("w:" + side)
                node.set(qn("w:w"), "90")
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            pr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                if ci == 1:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.bold = ri == 0
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if ri == 0 else "000000")
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t


para("Implementación del avance dos de ORVIX", "Title")
para("Interfaz de escritorio y distribución para Windows", "Subtitle")
para("Versión 0.2.1\nRevisión del 12 de octubre de 2026", "Normal")
para("La interfaz RAW ya está conectada al motor nativo y se distribuye mediante un instalador. Este anexo actualiza el estado de implementación del documento académico. Los diseños de filtros, detección e histórico SQLite se mantienen como trabajo de entregas posteriores.")
heading("Interfaz y conexión al motor")
para("La aplicación permite elegir una cámara o la demo sintética, iniciar y detener la captura y revisar el diagnóstico sin usar una terminal. La vista conserva la proporción del video; admite zoom con la rueda o los botones y desplazamiento al arrastrar una imagen ampliada. Las métricas diferencian FPS del productor, FPS del consumidor, saltos de secuencia, sobrescrituras y ocupación del anillo.")
para("Capture conserva la cámara en C++. QProcess gestiona su ciclo de vida; un QThread consume memoria compartida, interpreta los bytes con NumPy y decodifica con OpenCV. Qt muestra la imagen en el hilo gráfico. Un buzón de un solo frame evita acumular video si la interfaz se retrasa. No se usa cv2.VideoCapture ni se guardan frames automáticamente.")
heading("Validación de la entrega")
table(["Comprobación", "Resultado", "Condición"], [
    ("Pruebas nativas CTest", "16 de 16", "MSVC Release x64"),
    ("Pruebas Python unitarias y layout", "13 de 13", "Qt offscreen disponible"),
    ("Pruebas de integración", "5 de 5", "Productor C++ real y demo"),
    ("Demo sintética", "30 FPS", "Vista RAW y parada limpia"),
    ("Webcam GENERAL WEBCAM", "19.46 FPS", "1280×720 y NV12"),
    ("Sobrescrituras durante prueba física", "0", "877 frames publicados"),
])
caption = para("La prueba física se realizó desde el ejecutable empaquetado y mostró imagen en la interfaz. Al detener, el motor publicó el fin de captura y liberó la cámara. Los 30 FPS son un objetivo nominal, no el rendimiento medido de esta webcam. Estas pruebas no certifican todas las cámaras ni todos los sistemas operativos.")
caption.paragraph_format.space_before = Pt(7)

heading("Instalación y demostración", new_page=True)
para("El paquete se dirige a Windows 10 y 11 de 64 bits. Incluye Python, Qt, NumPy, OpenCV y el motor nativo; no requiere preparar un entorno de programación. El instalador crea la aplicación por usuario y permite agregar un acceso directo.")
for text in (
    "Ejecutar ORVIX-Avance2-Setup-0.2.1.exe, ubicado en la carpeta Aplicacion de los entregables del avance dos, y completar el instalador en español.",
    "Abrir ORVIX desde el menú Inicio. La aplicación no activa automáticamente la webcam.",
    "Dejar Demo sintética como fuente y pulsar Iniciar para demostrar el puente sin hardware. Comprobar la imagen RAW, la secuencia y las métricas.",
    "Pulsar Detener, seleccionar la webcam e iniciar de nuevo. Buscar actualiza los dispositivos disponibles sin abrirlos.",
    "Ampliar la vista con la rueda o con los botones. Arrastrar para moverse y usar Ajustar o doble clic para recuperar la vista completa.",
    "Abrir Diagnóstico para revisar los eventos. Guardar diagnóstico permite elegir un archivo de texto; no exporta video ni imágenes.",
    "Terminar con Detener. Si se cierra una ventana con captura activa, confirmar la parada para liberar los recursos antes de salir.",
):
    para(text, "List Number")
heading("Atajos y fallos controlados")
para("Ctrl+Enter inicia la captura; Esc la detiene; F11 alterna pantalla completa y Ctrl+0 ajusta la imagen. Los controles de selección quedan deshabilitados mientras se captura para evitar cambios de fuente inconsistentes.")
para("Si el consumidor deja de recibir frames durante cinco segundos, la interfaz muestra un aviso y deja de presentar una imagen antigua como si fuera video activo. Reconectar vista reabre el consumidor sin reiniciar Capture. Si se solicita detener y el motor no responde en tres segundos, se finaliza únicamente el proceso de esa captura. Cerrar o fallar el consumidor por sí solo no detiene al productor.")
heading("Alcance y archivos de entrega")
para("La instalación queda en %LOCALAPPDATA%/Programs/ORVIX. Los logs técnicos quedan en %LOCALAPPDATA%/ORVIX/ORVIX/logs y no se envían a Internet. La desinstalación conserva logs y preferencias. El paquete no tiene firma digital; no se deben desactivar controles de seguridad para ocultar esa condición.")
para("Esta entrega implementa captura nativa, Bridge y visualización RAW. La base de datos y el histórico del documento principal siguen siendo diseños; no hay filtros o detección implementados. El código, las pruebas y las recetas de empaquetado se versionan en https://github.com/Jeamlake/orvix. Los ejecutables se generan con scripts/package-desktop.ps1 y están excluidos de Git.")
para("Evidencia reproducible: docs/es/evidence/delivery-02/desktop-validation.md. Pruebas: scripts/test.ps1 y python -m pytest -m integration. Código de interfaz: python/src/orvix/ui/desktop.py, capture_controller.py y widgets.py.")
doc.save(target)
print(target)
